"""Kaynak okumasını zaman sınırı olan, anahtarsız alt süreçte çalıştır."""
import multiprocessing as mp
import time

_GOM_COCUK = {}  # alt surec ici vektor oturum onbellegi (suru basina bir)


def _perform(kind, args):
    if kind == "extract":
        from .extract import extract
        return extract(*args)
    if kind == 'extract_window':
        from .extract import extract_window
        return extract_window(*args)
    if kind == 'image_ocr':
        from .image_ocr import ocr_images
        return ocr_images(*args)
    if kind == 'image':
        from PIL import Image, ImageOps
        path, number = args
        with Image.open(path) as im:
            total=getattr(im,'n_frames',1); im.seek(max(0,min(number-1,total-1)))
            rgb=ImageOps.exif_transpose(im).convert('RGB'); rgb.thumbnail((1600,1600))
            return dict(number=number,total=total,data=rgb.tobytes(),width=rgb.width,height=rgb.height)
    if kind == "ocr":
        from .ocr import ocr_document
        return ocr_document(*args)
    if kind == "thumb":
        return _thumb(*args)
    if kind == "pdf":
        path, number, zoom, fit = (list(args) + [None])[:4]
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(path) as doc:
            number = max(1, min(number, len(doc)))
            page = doc[number - 1]
            width, height = page.get_size()
            if fit:
                # fit: (en fazla genişlik px, en fazla yükseklik px); None = sınır yok
                sinirlar = [x / y for x, y in zip(fit, (width, height)) if x]
                olcek = min(sinirlar) if sinirlar else 2.0 * zoom
            else:
                olcek = 2.0 * zoom
            scale = min(olcek, (8_000_000 / max(1, width * height)) ** 0.5)
            image = page.render(scale=scale).to_pil().convert("RGB")
            return {"number": number, "total": len(doc), "width": image.width,
                    "height": image.height, "data": image.tobytes(),
                    "sheet_width": width, "sheet_height": height, "scale": scale}
    if kind == "excel":
        path, sheet = args
        from .extract import _zip_guard
        if path.lower().endswith(".xlsx"):
            if _zip_guard(path):
                raise ValueError("ZIP_LIMIT")
            import openpyxl
            from itertools import islice
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
                rows = list(islice(ws.iter_rows(max_col=50, values_only=True), 200))
                return {"rows": rows, "sheet": ws.title}
            finally:
                wb.close()
        import xlrd
        wb = xlrd.open_workbook(path, on_demand=True)
        try:
            ws = wb.sheet_by_name(sheet) if sheet in wb.sheet_names() else wb.sheet_by_index(0)
            return {"rows": [ws.row_values(i, 0, min(50, ws.ncols))
                              for i in range(min(200, ws.nrows))], "sheet": ws.name}
        finally:
            wb.release_resources()
    if kind == "gom":
        from . import embed as embed_mod

        model_yolu, metinler = args
        ot = _GOM_COCUK.get(model_yolu)
        if ot is None:
            ot = embed_mod.Oturum(model_yolu)
            _GOM_COCUK[model_yolu] = ot
        return {"vektorler": ot.encode(metinler, gorev="belge").tolist()}
    raise ValueError("bilinmeyen iş")


def _thumb(kaynak, hedef, olcu, sayfa=1):
    """Kaynağın istenen sayfasını/kadrajını küçük resme indirgeyip dosyaya yazar.

    PDF sayfası pypdfium2 ile ~2x çözünürlükte render edilir, ardından
    Pillow ile hedef ölçüye küçültülür (küçültme, doğrudan düşük ölçekli
    render'dan daha okunur sonuç verir). Görsellerde EXIF yönü uygulanır;
    çok kareli görsellerde (TIFF) `sayfa` karesi seçilir.
    """
    from PIL import Image, ImageOps
    import os
    hedef_dizin = os.path.dirname(hedef)
    if hedef_dizin:
        os.makedirs(hedef_dizin, exist_ok=True)
    ext = os.path.splitext(kaynak)[1].lower()
    if ext == ".pdf":
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(kaynak) as doc:
            if len(doc) < 1:
                raise ValueError("EMPTY_PDF")
            page = doc[max(0, min(int(sayfa) - 1, len(doc) - 1))]
            width, height = page.get_size()
            scale = min(4.0, max(0.2, (olcu * 2.0) / max(1.0, max(width, height))))
            image = page.render(scale=scale).to_pil().convert("RGB")
    else:
        with Image.open(kaynak) as ham:
            total = getattr(ham, "n_frames", 1)
            if total > 1:
                ham.seek(max(0, min(int(sayfa) - 1, total - 1)))
            ham.draft("RGB", (olcu * 2, olcu * 2))
            image = ImageOps.exif_transpose(ham).convert("RGB")
    image.thumbnail((olcu, olcu), Image.LANCZOS)
    if hedef.lower().endswith(".webp"):
        image.save(hedef, "WEBP", quality=82, method=4)
    elif hedef.lower().endswith(".png"):
        image.save(hedef, "PNG", optimize=True)
    else:
        image.save(hedef, "JPEG", quality=85, optimize=True)
    return {"file": hedef, "width": image.width, "height": image.height}


def _child(sender, kind, args):
    try:
        import pickle
        payload = pickle.dumps((True, _perform(kind, args)))
        if len(payload) > 40 * 1024 * 1024:
            payload = pickle.dumps((False, "SIZE_LIMIT"))
        sender.send_bytes(payload)
    except Exception as error:
        import pickle
        sender.send_bytes(pickle.dumps((False, type(error).__name__)))
    finally:
        sender.close()


def run_work(kind, args, timeout=60, stop=None):
    import pickle
    context = mp.get_context("spawn")
    receiver, sender = context.Pipe(duplex=False)
    process = context.Process(target=_child, args=(sender, kind, args), daemon=True)
    process.start()
    sender.close()
    deadline = time.monotonic() + timeout
    try:
        while not receiver.poll(0.05):
            if stop and stop():
                raise InterruptedError("iş iptal edildi")
            if time.monotonic() >= deadline:
                raise TimeoutError("PROCESS_TIMEOUT")
            if not process.is_alive():
                raise RuntimeError("çalışan beklenmedik biçimde kapandı")
        ok, result = pickle.loads(receiver.recv_bytes(40 * 1024 * 1024))
        if not ok:
            raise RuntimeError(result)
        return result
    finally:
        receiver.close()
        _terminate(process)


def _terminate(process) -> None:
    """Alt süreci ağacıyla birlikte kapat (Windows: taskkill /T)."""
    if process.is_alive():
        import os
        if os.name == "nt":
            import subprocess
            subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                           creationflags=subprocess.CREATE_NO_WINDOW, timeout=10)
        process.terminate()
    process.join(3)
    if process.is_alive():
        process.kill()
        process.join()
    process.close()


def _serve(req_reader, resp_writer):
    """Kalıcı işçi döngüsü: (kind,args) isteği al, yürüt, sonucu geri yolla."""
    import pickle
    while True:
        try:
            kind, args = pickle.loads(req_reader.recv_bytes(40 * 1024 * 1024))
        except (EOFError, OSError):
            break
        try:
            payload = pickle.dumps((True, _perform(kind, args)))
            if len(payload) > 40 * 1024 * 1024:
                payload = pickle.dumps((False, "SIZE_LIMIT"))
        except Exception as error:
            payload = pickle.dumps((False, type(error).__name__))
        try:
            resp_writer.send_bytes(payload)
        except OSError:
            break
    resp_writer.close()
    req_reader.close()


class OcrSession:
    """Tarama turu boyunca yeniden kullanılan spawn işçisi.

    RapidOCR+ONNX yükleme maliyeti (~5-10 sn) dosya başına değil tur başına
    bir kez ödenir. Timeout/iptal/çökme sonrası süreç öldürülür; çağıran
    bir sonraki iş için yeni oturum kurar.
    """

    def __init__(self):
        import pickle
        self._pickle = pickle
        context = mp.get_context("spawn")
        # iste: ebeveyn yazar (req_w), cocuk okur (req_r); cevap: cocuk yazar, ebeveyn okur.
        req_r, req_w = context.Pipe(duplex=False)
        resp_r, resp_w = context.Pipe(duplex=False)
        self.process = context.Process(target=_serve, args=(req_r, resp_w), daemon=True)
        self.process.start()
        self._req_w = req_w
        self._resp_r = resp_r
        req_r.close()   # ebeveyn okuma ucunu tasimaz
        resp_w.close()  # ebeveyn cocugun gonderme ucunu tasimaz

    def alive(self) -> bool:
        return self.process is not None and self.process.is_alive()

    def call(self, kind, args, timeout=600, stop=None):
        if not self.alive():
            raise RuntimeError("çalışan beklenmedik biçimde kapandı")
        self._req_w.send_bytes(self._pickle.dumps((kind, args)))
        deadline = time.monotonic() + timeout
        while not self._resp_r.poll(0.05):
            if stop and stop():
                self.kill()
                raise InterruptedError("iş iptal edildi")
            if time.monotonic() >= deadline:
                self.kill()
                raise TimeoutError("PROCESS_TIMEOUT")
            if not self.process.is_alive():
                self.kill()
                raise RuntimeError("çalışan beklenmedik biçimde kapandı")
        ok, result = self._pickle.loads(self._resp_r.recv_bytes(40 * 1024 * 1024))
        if not ok:
            raise RuntimeError(result)
        return result

    def kill(self) -> None:
        if self.process is not None:
            _terminate(self.process)
            self.process = None
        for pipe in (self._req_w, self._resp_r):
            try:
                pipe.close()
            except OSError:
                pass


class WorkerSession(OcrSession):
    """Kalıcı genel işçi oturumu: tekrarlı kaynak işleri tek süreçte koşar.

    `OcrSession` ile aynı sözleşmeyi kullanır (`call(kind, args)`); adı işin
    türünü değil, yaşam döngüsünü anlatır (ör. thumbnail üretimi).
    """
