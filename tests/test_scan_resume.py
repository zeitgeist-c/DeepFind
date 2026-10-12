"""Resume/idempotency: ayni imza tekrar kuyruga girmez, kapsam kalicidir."""
import os
import sqlite3

from evrakci import coordinator as C
from evrakci import migrate, queue as q
from evrakci import schema
from evrakci.extract import ExtractResult, Segment
from evrakci.scan import Entry
from evrakci.vocab import ensure_v5


def database(path=None):
    con = sqlite3.connect(str(path) if path else ":memory:")
    schema.ensure_schema(con)
    ensure_v5(con)
    migrate.ensure_current(con)
    return con


def _entry(path, size=10, mtime_ns=1000):
    name = os.path.basename(path)
    return Entry(path, name, os.path.splitext(name)[1].lower(), size, mtime_ns)


def test_scan_selected_survives_reconnect(tmp_path):
    db = tmp_path / "kasa.db"
    con = database(db)
    con.execute("INSERT OR IGNORE INTO scan_selected VALUES (7)")
    con.commit()
    con.close()
    con2 = database(db)  # ensure_current tekrar kosar (idempotent olmali)
    try:
        assert con2.execute("SELECT count(*) FROM scan_selected").fetchone()[0] == 1
    finally:
        con2.close()


def test_same_signature_not_requeued(tmp_path):
    con = database(tmp_path / "kasa.db")
    try:
        entry = _entry("C:/arsiv/a.pdf")
        assert C.note_discovered(con, entry) == "queued"
        fid = con.execute("SELECT file_id FROM files WHERE path=?", (entry.path,)).fetchone()[0]
        sig = q.signature_of(entry.size, entry.mtime_ns)
        C.write_revision(con, fid, sig,
                         ExtractResult([Segment("pdf_page", 1, "Sayfa 1", "tahsilat makbuzu")]), "")
        jid = con.execute("SELECT job_id FROM jobs").fetchone()[0]
        q.complete(con, jid)
        assert C.note_discovered(con, entry) == "skipped"
        assert q.pending_count(con) == 0
    finally:
        con.close()


def test_interrupted_run_resumes_same_run_id(tmp_path):
    con = database(tmp_path / "kasa.db")
    try:
        migrate.ensure_archive(con)
        rid = con.execute(
            "INSERT INTO scan_runs(root, started_at) VALUES (?, ?)", ("/root", 1)).lastrowid
        con.execute("INSERT INTO scan_scopes(run_id, dir, complete) VALUES (?, ?, ?)",
                    (rid, "/root/a", 1))
        con.commit()
        assert C.find_resumable_run(con, "/root") == rid
        con.execute("UPDATE scan_runs SET finished_at=?, result=? WHERE run_id=?", (2, "complete", rid))
        con.commit()
        assert C.find_resumable_run(con, "/root") is None
    finally:
        con.close()


def test_manifest_run_never_resumed(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"%PDF-1.4 a")
    con = database(tmp_path / "kasa.db")
    try:
        rid, _total = C.build_manifest(con, str(tmp_path))
        absroot = os.path.abspath(str(tmp_path))
        assert C.find_resumable_run(con, absroot) is None
        con.execute("UPDATE scan_runs SET finished_at=NULL WHERE run_id=?", (rid,))
        con.commit()
        assert C.find_resumable_run(con, absroot) is None
    finally:
        con.close()


def test_batched_discovery_defers_durability(tmp_path):
    db = tmp_path / "kasa.db"
    con = database(db)
    entry = _entry("C:/arsiv/a.pdf")
    assert C.note_discovered(con, entry, commit=False) == "queued"
    assert q.pending_count(con) == 1  # ayni baglantida gorunur
    con.close()  # commit yok -> geri alinir
    con2 = database(db)
    try:
        assert con2.execute("SELECT count(*) FROM files").fetchone()[0] == 0
        assert C.note_discovered(con2, entry) == "queued"  # varsayilan kalici
    finally:
        con2.close()
    con3 = database(db)
    try:
        assert con3.execute("SELECT count(*) FROM files").fetchone()[0] == 1
    finally:
        con3.close()


def test_changed_signature_requeued(tmp_path):
    con = database(tmp_path / "kasa.db")
    try:
        entry = _entry("C:/arsiv/a.pdf", size=10, mtime_ns=1000)
        assert C.note_discovered(con, entry) == "queued"
        fid = con.execute("SELECT file_id FROM files WHERE path=?", (entry.path,)).fetchone()[0]
        sig = q.signature_of(entry.size, entry.mtime_ns)
        C.write_revision(con, fid, sig,
                         ExtractResult([Segment("pdf_page", 1, "Sayfa 1", "tahsilat makbuzu")]), "")
        jid = con.execute("SELECT job_id FROM jobs").fetchone()[0]
        q.complete(con, jid)
        changed = _entry("C:/arsiv/a.pdf", size=11, mtime_ns=1000)
        assert C.note_discovered(con, changed) == "queued"
        assert q.pending_count(con) == 1
    finally:
        con.close()


def test_completed_dirs_skipped_on_resume(tmp_path):
    from openpyxl import Workbook
    import time as _time
    root = tmp_path / "kaynak"
    root.mkdir()
    done = root / "bitmis"
    done.mkdir()
    todo = root / "yeni"
    todo.mkdir()
    for d, name in ((done, "a.xlsx"), (todo, "b.xlsx")):
        wb = Workbook()
        wb.active.append(["kayit"])
        wb.save(d / name)
    past = _time.time() - 10
    for p in (done / "a.xlsx", todo / "b.xlsx", done):
        os.utime(p, (past, past))
    con = database(tmp_path / "kasa.db")
    try:
        absroot = os.path.abspath(str(root))
        now = int(_time.time())
        rid = con.execute(
            "INSERT INTO scan_runs(root, started_at) VALUES (?, ?)", (absroot, now)).lastrowid
        con.execute("INSERT INTO scan_scopes(run_id, dir, complete) VALUES (?, ?, 1)",
                    (rid, os.path.abspath(str(done))))
        con.commit()
        stats = C.run_scan(con, str(tmp_path / "kasa.db"), str(root))
        assert stats["found"] == 1
        wb = Workbook()
        wb.active.append(["yeni"])
        wb.save(done / "c.xlsx")
        os.utime(done / "c.xlsx", (past, past))
        stats2 = C.run_scan(con, str(tmp_path / "kasa.db"), str(root))
        assert stats2["found"] >= 2
    finally:
        con.close()


def test_progress_distinguishes_discovery_and_work():
    events = []
    C.emit_progress(events.append, stage="Klasörler keşfediliyor",
                    found=100, done=10, pending=90)
    assert events[0]["phase"] == "kesif"
    assert events[0]["found"] == 100
    C.emit_progress(events.append, stage="İçerik okunuyor",
                    found=100, done=10, pending=90)
    assert events[1]["phase"] == "islem"


def test_stop_flag_cancels_scan(tmp_path):
    from openpyxl import Workbook
    import time as _time
    root = tmp_path / "kaynak"
    root.mkdir()
    wb = Workbook()
    wb.active.append(["kayit"])
    wb.save(root / "a.xlsx")
    past = _time.time() - 10
    os.utime(root / "a.xlsx", (past, past))
    con = database(tmp_path / "kasa.db")
    try:
        stats = C.run_scan(con, str(tmp_path / "kasa.db"), str(root), stop=lambda: True)
        assert stats["result"] == "cancelled"
        assert con.execute(
            "SELECT result FROM scan_runs").fetchone()[0] == "cancelled"
    finally:
        con.close()


def test_run_scan_resumes_unfinished_run(tmp_path):
    root = tmp_path / "kaynak"
    root.mkdir()
    (root / "a.xlsx").write_bytes(b"degil")
    con = database(tmp_path / "kasa.db")
    try:
        absroot = os.path.abspath(str(root))
        rid = con.execute(
            "INSERT INTO scan_runs(root, started_at) VALUES (?, ?)", (absroot, 1)).lastrowid
        con.commit()
        stats = C.run_scan(con, str(tmp_path / "kasa.db"), str(root))
        rows = con.execute("SELECT run_id FROM scan_runs").fetchall()
        assert [r[0] for r in rows] == [rid]
        assert con.execute(
            "SELECT finished_at FROM scan_runs WHERE run_id=?", (rid,)).fetchone()[0] is not None
        assert stats["found"] >= 1
    finally:
        con.close()
