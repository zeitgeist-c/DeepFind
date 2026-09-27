"""A3 dikey akis testleri: sentetik belge -> tara -> indeksle -> ara -> snippet."""
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from evrakci import extract, index, scan, schema, search, vault  # noqa: E402
from evrakci.migrate import ensure_v6  # noqa: E402
from evrakci.vocab import ensure_v5  # noqa: E402

PW = "test-sifresi-12345"


def _sentetik(kok: str) -> None:
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("TR", "C:/Windows/Fonts/arial.ttf"))
    c = canvas.Canvas(os.path.join(kok, "Ankara Vergi Levhası 2024.pdf"))
    c.setFont("TR", 14)
    c.drawString(72, 750, "Ankara Ayşe 2024 vergi levhası")
    c.showPage()
    c.setFont("TR", 14)
    c.drawString(72, 750, "ikinci sayfada beyanname detayı")
    c.save()

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vergi"
    ws["B14"] = "Ayşe Yılmaz"
    ws["C14"] = "Ankara"
    wb.save(os.path.join(kok, "liste 2024.xlsx"))

    from docx import Document
    d = Document()
    d.add_paragraph("Mehmet Demir Aralık faturası İstanbul")
    d.save(os.path.join(kok, "not.docx"))

    with open(os.path.join(kok, "eski-arşiv.doc"), "wb") as f:
        f.write(b"\x00\x01sahte-doc")


def test_kasa_yanlis_sifre_acmaz(tmp_path):
    db = str(tmp_path / "k.db")
    con = vault.open_vault(db, PW)
    schema.ensure_schema(con)
    ensure_v5(con)
    ensure_v6(con)
    con.close()
    with pytest.raises(Exception):
        vault.open_vault(db, "yanlis-sifre-00000")


def test_dikey_akis(tmp_path):
    kok = tmp_path / "arsiv"
    kok.mkdir()
    _sentetik(str(kok))
    db = str(tmp_path / "k.db")
    con = vault.open_vault(db, PW)
    schema.ensure_schema(con)
    ensure_v5(con)
    ensure_v6(con)
    girisler = list(scan.walk_ro(str(kok)))
    assert len(girisler) == 4
    for e in girisler:
        index.index_entry(con, e, extract.extract(e.path, e.ext, e.size))
    # sorgu 1: iki kelime ayni belgede
    r1 = search.search(con, "Ankara Ayşe")
    assert r1 and r1[0]["name"].endswith(".pdf"), r1
    assert "Ankara" in r1[0]["snippet"] and "Ayşe" in r1[0]["snippet"]
    assert r1[0]["page"] == 1
    # sorgu 2: xlsx hucre
    r2 = search.search(con, "Ayşe Yılmaz")
    assert any(x["ext"] == ".xlsx" for x in r2), r2
    # sorgu 3: doc ad/yol
    r3 = search.search(con, "eski-arşiv")
    assert r3 and r3[0]["ext"] == ".doc", r3
    # guvenli sorgu: FTS enjeksiyonu calismaz, hata da vermez
    r4 = search.search(con, 'vergi" OR 1=1 --')
    assert isinstance(r4, list)
    con.close()


def test_surum_esitligi():
    import tomllib
    from evrakci.version import __version__
    with open(os.path.join(os.path.dirname(__file__), "..", "pyproject.toml"), "rb") as f:
        project = tomllib.load(f)
        assert project["tool"]["setuptools"]["dynamic"]["version"]["attr"] == "evrakci.version.__version__"
        assert "version" in project["project"]["dynamic"]


def test_router_bayrak_degerini_komut_sanmaz():
    from evrakci.__main__ import _ilk_komut, main
    assert _ilk_komut(["--kasa", "C:/x/k.db", "--sifre", "pw", "tani"]) == "tani"
    assert _ilk_komut(["--kasa", "C:/x/k.db"]) is None
    assert _ilk_komut([]) is None
    assert _ilk_komut(["tara", "C:/arsiv"]) == "tara"
    with pytest.raises(SystemExit) as exc:
        main(["--help"])
    assert exc.value.code == 0
