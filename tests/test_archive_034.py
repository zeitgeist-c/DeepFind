"""Archive completeness, retrieval correctness and bounded UI regressions."""
import sqlite3,time,os
import pytest
from evrakci import schema,migrate,index,rank,coordinator,extract,vault
from evrakci.scan import Entry
from evrakci.extract import Segment,ExtractResult
from evrakci.vocab import ensure_v5

def database():
    con=sqlite3.connect(':memory:'); schema.ensure_schema(con); ensure_v5(con); migrate.ensure_current(con); return con

def test_large_archive_does_not_lose_intersection_or_filtered_matches():
    con=database()
    for n in range(510):
        entry=Entry(f'C:/archive/{n:04}.pdf',f'{n}.pdf','.pdf',1,1)
        index.index_entry(con,entry,ExtractResult([Segment('pdf_page',1,'Sayfa 1','tahsilat makbuzu' if n==509 else 'tahsilat')]))
    result=rank.search_ranked(con,'tahsilat makbuzu')['results']
    assert len(result)==1 and result[0]['name']=='509.pdf'
    assert rank.search_ranked(con,'tahsilat',folder='0509')['results'][0]['name']=='509.pdf'
    assert rank.search_ranked(con,'tahsilat',offset=500)['results']
    con.close()


def test_phrase_query_and_identifiers():
    con=database()
    for n,text in enumerate(['tahsilat makbuzu 54321 18.750 TL','makbuzu önce tahsilat sonra 54321']):
        index.index_entry(con,Entry(f'{n}.pdf',f'{n}.pdf','.pdf',1,1),ExtractResult([Segment('pdf_page',1,'Sayfa 1',text)]))
    assert [r['name'] for r in rank.search_ranked(con,'"tahsilat makbuzu"')['results']]==['0.pdf']
    assert rank.search_ranked(con,'54321 18.750')['results'][0]['name']=='0.pdf'
    con.close()


def test_excel_windows_keep_last_row_and_source_unchanged(tmp_path):
    from openpyxl import Workbook
    file=tmp_path/'random.xlsx'; book=Workbook()
    for i in range(1100): book.active.append([f'kayıt {i}'])
    book.save(file); original=file.read_bytes()
    first=extract.extract_window(str(file),'.xlsx',len(original))
    assert first.next_cursor==1000
    last=extract.extract_window(str(file),'.xlsx',len(original),first.next_cursor)
    assert last.next_cursor is None and '1099' in last.segments[-1].raw_text
    assert file.read_bytes()==original

def test_reindex_failure_preserves_old_content(tmp_path):
    root=tmp_path/'source'; root.mkdir(); file=root/'file.xlsx'; file.write_bytes(b'bad')
    past=time.time()-10; os.utime(file,(past,past))
    con=database(); signature=f'{file.stat().st_size}:{file.stat().st_mtime_ns}'
    fid=index.index_entry(con,Entry(str(file),file.name,'.xlsx',3,1),ExtractResult([Segment('sheet',None,'sheet','eski tahsilat içeriği')]))
    from evrakci import job_queue as queue
    queue.enqueue(con,fid,signature); con.commit(); task=queue.claim(con)
    outcome=coordinator._process_job(con,task,[],str(tmp_path),str(root))
    assert outcome=='retry'
    assert 'eski tahsilat' in con.execute("SELECT raw_text FROM segments WHERE kind='sheet'").fetchone()[0]
    assert con.execute('SELECT index_status FROM files').fetchone()[0]=='partial'
    con.close()

def test_chunk_offsets_cover_entire_source():
    original='abc '*4000; parts=list(extract.chunks(Segment('para',None,'P',original)))
    assert ''.join(part.raw_text for part,_,_ in parts)==original
    assert parts[-1][2]==len(original)

def test_image_missing_ocr_stays_partial_and_queued(tmp_path,monkeypatch):
    from PIL import Image
    from evrakci import ocr
    root=tmp_path/'images'; root.mkdir(); image=root/'receipt.png'; Image.new('RGB',(60,60),'white').save(image)
    past=time.time()-10; os.utime(image,(past,past))
    db=str(tmp_path/'index.db'); con=vault.open_vault(db,'synthetic-archive-1234'); migrate.ensure_current(con)
    monkeypatch.setattr(ocr,'engine_ok',lambda *args:(False,['rapidocr-zinciri']))
    result=coordinator.run_scan(con,db,str(root))
    assert not result['complete']
    assert con.execute('SELECT index_status FROM files').fetchone()[0]=='partial'
    assert con.execute("SELECT count(*) FROM jobs WHERE kind='ocr'").fetchone()[0]==1
    con.close()

def test_windowed_extraction_can_resume_after_cancellation(tmp_path,monkeypatch):
    from openpyxl import Workbook
    root=tmp_path/'archive'; root.mkdir(); file=root/'long.xlsx'; book=Workbook()
    for i in range(1100): book.active.append([f'evrak {i}'])
    book.save(file); past=time.time()-10; os.utime(file,(past,past))
    db=str(tmp_path/'index.db'); con=vault.open_vault(db,'synthetic-archive-1234'); migrate.ensure_current(con)
    original=coordinator._process_job; stopped=False
    def process(*args,**kwargs):
        nonlocal stopped
        result=original(*args,**kwargs); stopped=True; return result
    monkeypatch.setattr(coordinator,'_process_job',process)
    coordinator.run_scan(con,db,str(root),stop=lambda:stopped)
    assert con.execute('SELECT cursor FROM jobs').fetchone()[0]==1000
    monkeypatch.setattr(coordinator,'_process_job',original)
    result=coordinator.run_scan(con,db,str(root))
    assert result['complete']
    assert rank.search_ranked(con,'1099')['results']
    con.close()

def test_pdf_windows_include_page_after_legacy_limit(tmp_path):
    import pypdfium2 as pdfium
    path=tmp_path/'long.pdf'; pdf=pdfium.PdfDocument.new()
    for _ in range(1002): pdf.new_page(100,100)
    pdf.save(path); pdf.close()
    result=extract.extract_window(str(path),'.pdf',path.stat().st_size,1000)
    assert [s.page for s in result.segments]==[1001,1002]
    assert result.next_cursor is None

def test_docx_windows_preserve_last_paragraph_and_offsets(tmp_path):
    from docx import Document
    path=tmp_path/'long.docx'; doc=Document()
    for i in range(105): doc.add_paragraph(f'evrak {i}')
    doc.save(path)
    first=extract.extract_window(str(path),'.docx',path.stat().st_size)
    assert first.next_cursor==100
    last=extract.extract_window(str(path),'.docx',path.stat().st_size,100)
    assert last.segments[-1].raw_text=='evrak 104' and last.next_cursor is None

def test_preview_loads_beyond_sixty_sections_and_discards_closed_result(tmp_path):
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    from PySide6.QtWidgets import QApplication
    from evrakci.preview import Onizleme
    app=QApplication.instance() or QApplication([])
    db=str(tmp_path/'preview.db'); password='synthetic-preview-1234'
    con=vault.open_vault(db,password); migrate.ensure_current(con)
    index.index_entry(con,Entry('C:/not-existing/long.docx','long.docx','.docx',1,1),ExtractResult([Segment('para',None,f'P {i}',f'Belge bölümü {i}') for i in range(65)]))
    preview=Onizleme(); result=dict(path='C:/not-existing/long.docx',name='long.docx',ext='.docx')
    def wait():
        deadline=time.monotonic()+10
        while preview.busy():
            app.processEvents(); time.sleep(.01); assert time.monotonic()<deadline
        app.processEvents()
    preview.goster(con,result,password); wait()
    for _ in range(3): preview._load_more(); wait()
    assert 'Belge bölümü 64' in preview._fallback
    preview.goster(con,result,password); preview.temizle(); wait()
    assert preview._path is None and not preview.metin.toPlainText()
    preview.close(); con.close()

def test_failed_migration_restores_verified_previous_vault(tmp_path,monkeypatch):
    os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
    from PySide6.QtWidgets import QApplication
    from evrakci.app import AnaPencere
    app=QApplication.instance() or QApplication([])
    db=str(tmp_path/'old.db'); password='synthetic-upgrade-1234'
    con=vault.open_vault(db,password); migrate.ensure_current(con)
    index.index_entry(con,Entry('old.pdf','old.pdf','.pdf',1,1),ExtractResult([Segment('pdf_page',1,'Sayfa 1','tahsilat korunmalı')]))
    con.execute("DELETE FROM schema_meta WHERE k='archive_version'"); con.commit(); con.close()
    original=migrate.ensure_current
    def failed(con):
        con.execute('DELETE FROM segments'); con.commit(); raise RuntimeError('synthetic migration failure')
    monkeypatch.setattr(migrate,'ensure_current',failed)
    with pytest.raises(RuntimeError,match='synthetic migration failure'): AnaPencere(db,password)
    monkeypatch.setattr(migrate,'ensure_current',original)
    con=vault.open_vault(db,password)
    assert rank.search_ranked(con,'tahsilat')['results']; con.close()
