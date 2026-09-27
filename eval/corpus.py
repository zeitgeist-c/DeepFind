"""Sentetik kalite corpus ureteci (yapay veri; gercek arsiv yok). Deterministik."""
from __future__ import annotations

import json
import os
import shutil

SEHIR = ["Ankara", "İstanbul", "İzmir"]
KISI = ["Ayşe Yılmaz", "Mehmet Demir", "Fatma Kaya", "Mustafa Çelik", "Elif Şahin",
        "Hasan Öztürk", "Zeynep Arslan", "Ali Vural"]
TUR = ["vergi levhası", "fatura", "beyanname", "sözleşme", "rapor", "bordro"]
YIL = ["2021", "2022", "2023", "2024"]


def _pdf(path, lines):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        pdfmetrics.getFont("TR")
    except KeyError:
        pdfmetrics.registerFont(TTFont("TR", "C:/Windows/Fonts/arial.ttf"))
    c = canvas.Canvas(path)
    c.setFont("TR", 13)
    for i, ln in enumerate(lines):
        if i and i % 8 == 0:
            c.showPage()
            c.setFont("TR", 13)
        c.drawString(72, 750 - (i % 8) * 30, ln)
    c.save()


def build(root: str) -> list:
    """Corpus kur + kayit listesi don: [{name, sehir, kisi, yil, tur, no, path}]."""
    if os.path.exists(root):
        shutil.rmtree(root)
    recs = []
    idx = 0
    for si, sehir in enumerate(SEHIR):
        for yi, yil in enumerate(YIL):
            for ti, tur in enumerate(TUR):
                kisi = KISI[(si + yi + ti) % len(KISI)]
                ilk = kisi.split()[0]
                d = os.path.join(root, sehir, yil)
                os.makedirs(d, exist_ok=True)
                ad = f"{ilk} {tur} {yil}"
                no = f"BV-{1000 + idx}"
                if idx % 3 == 0:
                    p = os.path.join(d, ad + ".pdf")
                    _pdf(p, [f"{sehir} {kisi} {yil} {tur}",
                             f"Belge no {no} duzenlendi",
                             "Sayfa iki detayi burada yer alir ek bilgi"])
                elif idx % 3 == 1:
                    import openpyxl
                    p = os.path.join(d, ad + ".xlsx")
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Belge"
                    ws["A1"] = sehir
                    ws["B2"] = kisi
                    ws["C3"] = f"{tur} {yil} {no}"
                    wb.save(p)
                    wb.close()
                else:
                    from docx import Document
                    p = os.path.join(d, ad + ".docx")
                    doc = Document()
                    doc.add_paragraph(f"{kisi} {sehir} {tur}")
                    doc.add_paragraph(f"{yil} yili {no} numarali rapor")
                    doc.save(p)
                recs.append({"name": os.path.basename(p), "sehir": sehir, "kisi": kisi,
                             "yil": yil, "tur": tur, "no": no, "path": p})
                idx += 1
    d = os.path.join(root, "Ankara", "2024")
    _pdf(os.path.join(d, "Daginik Belge 2024.pdf"),
         ["Ayşe Yılmaz burada"] + ["lorem ipsum dolor sit amet"] * 7 +
         ["vergi levhası ektedir son"])
    recs.append({"name": "Daginik Belge 2024.pdf", "sehir": "Ankara", "kisi": "Ayşe Yılmaz",
                 "yil": "2024", "tur": "vergi levhası", "no": "BV-9000",
                 "path": os.path.join(d, "Daginik Belge 2024.pdf")})
    shutil.copy(os.path.join(d, "Daginik Belge 2024.pdf"),
                os.path.join(d, "Daginik Kopya 2024.pdf"))
    recs.append({"name": "Daginik Kopya 2024.pdf", "sehir": "Ankara", "kisi": "Ayşe Yılmaz",
                 "yil": "2024", "tur": "vergi levhası", "no": "BV-9000",
                 "path": os.path.join(d, "Daginik Kopya 2024.pdf"), "kopya": True})
    bas = "X" * 70000
    for n, son in (("benzer1.pdf", "BITIS-A"), ("benzer2.pdf", "BITIS-B")):
        with open(os.path.join(root, n), "wb") as f:
            f.write(b"%PDF " + bas.encode() + b" " + son.encode())
    open(os.path.join(root, "bos.pdf"), "wb").close()
    with open(os.path.join(root, "bozuk.pdf"), "wb") as f:
        f.write(b"bozuk icerik \xff\xfe")
    return recs


def make_queries(recs: list) -> list:
    """(sorgu, beklenen_ad, grup). Beklenenler uretim kayitlarindan alinir."""
    by = {(r["sehir"], r["kisi"].split()[0], r["yil"], r["tur"]): r["name"] for r in recs
          if not r.get("kopya")}
    ilk = recs[0]
    xlsx = next(r for r in recs if r["name"].endswith(".xlsx"))
    ilk_levha = next(r for r in recs if r["tur"] == "vergi levhası" and not r.get("kopya"))
    q = []
    q.append((f"{ilk['sehir']} {ilk['kisi']} {ilk['yil']} {ilk['tur']}", ilk["name"], "tam"))
    q.append((f"{xlsx['sehir']} {xlsx['kisi'].split()[0]} {xlsx['yil']} {xlsx['tur']}",
              xlsx["name"], "kisi"))
    q.append((f"Ayse vergi levhasi {ilk_levha['yil']}", ilk_levha["name"], "diakritik"))
    izm = next(r for r in recs if r["sehir"] == "İzmir" and r["tur"] == "rapor")
    q.append((f"Izmir rapor {izm['yil']}", izm["name"], "i-donusumu"))
    q.append((f"vergile evha {ilk_levha['yil']}", ilk_levha["name"], "typo"))
    q.append(("beyaname 2022", "YILSET:2022:beyanname", "typo-acik"))
    q.append(("Ankara Ayşe vergi", "Daginik Belge 2024.pdf", "capraz"))
    q.append((f'"{ilk_levha["tur"]}"', ilk_levha["name"], "tirnak"))
    q.append(("zebra helikopter kuantum", None, "no-result"))
    q.append(("1999", None, "no-result"))
    q.append(("", None, "bos"))
    for (sehir, ad, yil, tur), name in sorted(by.items()):
        q.append((f"{sehir} {ad} {yil} {tur}", name, "uretim"))
    return q


if __name__ == "__main__":
    import sys
    root = sys.argv[1]
    recs = build(root)
    qs = make_queries(recs)
    with open(os.path.join(os.path.dirname(root), "eval_queries.json"), "w",
              encoding="utf-8") as f:
        json.dump(qs, f, ensure_ascii=False, indent=1)
    print(f"belge={len(recs)} sorgu={len(qs)}")
